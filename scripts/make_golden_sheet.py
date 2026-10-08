"""Build the golden-50 labeling spreadsheet (dropdowns + guide tab). No AI involved.

Usage:
    uv run --with openpyxl python scripts/make_golden_sheet.py
"""

import csv
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

SRC = Path("data/raw/golden_50_to_label.csv")
OUT = Path("evals/golden_50_labeling.xlsx")

TOPICS = ["access", "usability", "playback", "downloads", "catalog", "billing", "support", "other"]
INTENTS = ["cancellation", "complaint", "request", "praise", "unclear"]
SEVERITY = ["1", "2", "3", "4", "5"]
SENTIMENT = ["-1", "-0.5", "0", "0.5", "1"]
YESNO = ["no", "yes"]

# Columns shown to the labeler: source info first, then the labels to fill in.
COLUMNS = [
    ("#", 5), ("review_text", 70), ("stars", 7),
    ("intent", 14), ("topic", 13), ("severity", 9), ("sentiment", 10),
    ("needs_review", 13), ("evidence_quote", 30), ("notes", 30), ("review_id", 38),
]

GUIDE = [
    ("HOW TO LABEL", ""),
    ("1. Read only the text.", "Ignore the stars: a 1-star 'Good' is still praise."),
    ("2. Intent (first that fits)", "cancellation > complaint > request > praise > unclear"),
    ("  cancellation", "Leaving, uninstalling, cancelling, or threatening to"),
    ("  complaint", "Negative experience, incl. mixed praise+criticism or generic 'bad app'"),
    ("  request", "Wants a change, nothing reported broken"),
    ("  praise", "Only positive"),
    ("  unclear", "Meaningless, unrelated, or bare boycott slogan"),
    ("3. Topic (pick one)", "Most severe problem; on a tie, the first mentioned. General praise = other"),
    ("  access", "Login, signup, password, account access"),
    ("  usability", "Navigation, controls, layout, queue/playlists, ADS"),
    ("  playback", "Won't play, crashes, lag, connection, audio quality, battery/data"),
    ("  downloads", "Downloading, offline/saved music, downloads disappearing"),
    ("  catalog", "Missing songs/artists, search, recommendations, lyrics"),
    ("  billing", "Price, charges, subscriptions, paywalls, Premium-only features"),
    ("  support", "Contacting support and their response"),
    ("  other", "General praise/criticism, unrelated, nothing specific"),
    ("  TRAP", "Mentioning Premium alone is NOT billing ('I pay and it crashes' = playback)"),
    ("4. Severity", ""),
    ("  1", "No problem: praise, unclear, or pure feature request"),
    ("  2", "Dislike / generic criticism / annoyance; nothing actually lost"),
    ("  3", "Something works worse or is restricted, but still usable"),
    ("  4", "A core task is blocked (can't log in, nothing plays)"),
    ("  5", "Explicit serious money, privacy or data harm"),
    ("  TRAP", "Anger, emojis or 'I'm cancelling' alone do NOT raise severity"),
    ("5. Sentiment", "-1 very negative, -0.5 negative, 0 neutral/mixed, 0.5 positive, 1 very positive"),
    ("6. needs_review", "'yes' if you honestly can't decide - write why in notes"),
    ("7. evidence_quote", "Optional: exact words from the review that justify your label"),
    ("", ""),
    ("Full guide with examples:", "docs/LABEL_GUIDE.md"),
]


def main():
    with SRC.open(encoding="utf-8-sig", newline="") as f:
        rows = list(csv.DictReader(f))
    assert len(rows) == 50, f"expected 50 golden rows, got {len(rows)}"

    wb = Workbook()
    ws = wb.active
    ws.title = "Label here"
    header_fill = PatternFill("solid", fgColor="1F4E79")
    fill_me = PatternFill("solid", fgColor="FFF2CC")

    for col, (name, width) in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=1, column=col, value=name)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        ws.column_dimensions[cell.column_letter].width = width
    ws.freeze_panes = "D2"

    for i, r in enumerate(rows, start=1):
        values = [i, r["review_text"], int(r["review_rating"]), "", "", "", "", "", "", "", r["review_id"]]
        for col, v in enumerate(values, start=1):
            cell = ws.cell(row=i + 1, column=col, value=v)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            if 4 <= col <= 8:
                cell.fill = fill_me

    last = len(rows) + 1
    for letter, options in [("D", INTENTS), ("E", TOPICS), ("F", SEVERITY), ("G", SENTIMENT), ("H", YESNO)]:
        dv = DataValidation(type="list", formula1='"' + ",".join(options) + '"', allow_blank=True,
                            showErrorMessage=True, error="Pick a value from the list")
        ws.add_data_validation(dv)
        dv.add(f"{letter}2:{letter}{last}")

    guide = wb.create_sheet("Guide")
    guide.column_dimensions["A"].width = 28
    guide.column_dimensions["B"].width = 90
    for i, (a, b) in enumerate(GUIDE, start=1):
        guide.cell(row=i, column=1, value=a).font = Font(bold=not a.startswith("  "))
        guide.cell(row=i, column=2, value=b)

    OUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT)
    print(f"Saved {OUT} with {len(rows)} reviews")


if __name__ == "__main__":
    main()
